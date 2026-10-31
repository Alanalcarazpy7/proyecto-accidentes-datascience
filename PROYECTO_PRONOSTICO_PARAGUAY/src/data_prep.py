"""Fase 1: reconstruccion desde XLSX oficiales, controles y trazabilidad."""
import json
import re

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from common import ROOT, TYPES, save_json, save_table, sha256

MONTHS = dict(enero=1, febrero=2, marzo=3, abril=4, mayo=5, junio=6,
              julio=7, agosto=8, setiembre=9, septiembre=9, octubre=10,
              noviembre=11, diciembre=12)


def main():
    raw = ROOT / 'data/raw'
    catalog = {r['archivo']: r for r in json.loads((raw / 'catalogo_fuentes.json').read_text(encoding='utf-8'))}
    records, annual, inventory, original, conversions = [], [], [], [], 0
    for file in sorted(raw.glob('*tipologia*.xlsx')):
        digest = sha256(file)
        if digest != catalog[file.name]['sha256']:
            raise ValueError(f'El hash original cambio: {file.name}')
        book = load_workbook(file, read_only=True, data_only=True)
        sheet = book.worksheets[0]
        year = None
        counts = 0
        for rownum, cells in enumerate(sheet.iter_rows(values_only=True), 1):
            label = str(cells[1] or '').strip().lower() if len(cells) > 1 else ''
            match = re.fullmatch(r'año\s+(20\d{2})', label)
            if match:
                year = int(match.group(1))
                annual.append({'ano': year, 'total_publicado': int(cells[2]), 'archivo': file.name,
                               'hoja': sheet.title, 'celda': f'C{rownum}'})
            elif year and label in MONTHS:
                row = {'fecha': f'{year}-{MONTHS[label]:02d}-01', 'accidentes': cells[2]}
                for i, name in enumerate(TYPES):
                    value = cells[i + 3]
                    original.append({'archivo': file.name, 'hoja': sheet.title,
                                     'celda': f'{chr(68+i)}{rownum}', 'campo': name,
                                     'valor_original': value})
                    if value == '-':
                        conversions += 1
                        value = 0
                    row[name] = value
                for key in ['accidentes'] + TYPES:
                    value = row[key]
                    if not isinstance(value, (float, int)) or not np.isfinite(value) or value < 0 or int(value) != value:
                        raise ValueError(f'Conteo ausente o invalido {file.name}:{rownum}:{key}={value!r}')
                    row[key] = int(value)
                if row['accidentes'] != sum(row[key] for key in TYPES):
                    raise ValueError(f'Las partes no suman el total: {file.name}:{rownum}')
                row.update(archivo=file.name, hoja=sheet.title, celda_total=f'C{rownum}',
                           celdas_tipos=f'D{rownum}:J{rownum}', edicion=int(file.name.split('_')[1]))
                records.append(row)
                counts += 1
        book.close()
        inventory.append({'archivo': file.name, 'formato': 'XLSX', 'bytes': file.stat().st_size,
                          'filas_mensuales': counts, 'sha256': digest, 'url': catalog[file.name]['url'],
                          'fecha_descarga_sesion': '2026-10-31', 'unidad': 'siniestros/mes',
                          'licencia': 'Acceso publico; licencia especifica no indicada en el XLSX'})
    observed = pd.DataFrame(records)
    save_table('filas_extraidas_con_versiones.csv', observed)
    save_table('celdas_originales_tipos.csv', pd.DataFrame(original))
    numeric = ['accidentes'] + TYPES
    conflicts = observed.groupby('fecha')[numeric].nunique().max(axis=1)
    if (conflicts > 1).any():
        raise ValueError('Hay versiones contradictorias: deben resolverse antes de modelar.')
    frame = observed.sort_values(['fecha', 'edicion']).drop_duplicates('fecha', keep='last').copy()
    frame['fecha'] = pd.to_datetime(frame['fecha'])
    expected = pd.date_range('2016-01-01', '2024-12-01', freq='MS')
    if not frame['fecha'].reset_index(drop=True).equals(pd.Series(expected)):
        raise ValueError('Faltan meses o hay periodos inesperados.')
    frame['ano'] = frame['fecha'].dt.year
    frame['mes'] = frame['fecha'].dt.month
    frame['semestre'] = np.where(frame['mes'] <= 6, 'Primero', 'Segundo')
    frame['dias_mes'] = frame['fecha'].dt.days_in_month
    frame['tasa_diaria_mes'] = frame['accidentes'] / frame['dias_mes']
    frame['particion'] = np.select([frame['ano'] <= 2022, frame['ano'] == 2023],
                                   ['Desarrollo', 'Calibracion'], default='Prueba')
    frame['fecha'] = frame['fecha'].dt.strftime('%Y-%m-%d')
    totals = frame.groupby('ano')['accidentes'].sum()
    for row in annual:
        if int(totals.loc[row['ano']]) != row['total_publicado']:
            raise ValueError(f'Total anual discrepante: {row}')
    target = ROOT / 'data/processed'
    target.mkdir(exist_ok=True)
    frame.to_csv(target / 'serie_mensual.csv', index=False, encoding='utf-8-sig')
    frame.loc[frame['particion'] == 'Desarrollo'].to_csv(target / 'desarrollo_2016_2022.csv', index=False, encoding='utf-8-sig')
    save_table('inventario_fuentes.csv', pd.DataFrame(inventory))
    save_table('controles_anuales.csv', pd.DataFrame(annual))
    save_table('totales_anuales.csv', totals.rename('siniestros').reset_index())
    bitacora = pd.DataFrame([
        ['Seleccion', len(inventory), 'Siete cuadros de tipologia con mismo alcance publicado; resto no se suma.'],
        ['Extraccion', len(observed), 'Filas mensuales; se excluyen titulos, fuentes y totales anuales.'],
        ['Guiones a cero', conversions, 'Solo en componentes categoricos; las sumas por fila y ano se conservan. No se imputan meses.'],
        ['Versiones repetidas', len(observed)-len(frame), 'Coinciden todos los conteos; se conserva la edicion mas reciente y se guardan todas las versiones.'],
        ['Normalizar fecha', len(frame), 'Primer dia representa el mes, no una fecha de accidente individual.'],
        ['Atipicos', 0, 'No se eliminan ni winsorizan meses validos; los cambios pueden ser reales.'],
        ['Imputar objetivo', 0, 'No hay faltantes ni meses inventados.'],
        ['Separar tiempo', len(frame), '84 desarrollo; 12 calibracion; 12 prueba.'],
    ], columns=['operacion', 'cantidad_afectada', 'fundamento'])
    save_table('bitacora_limpieza.csv', bitacora)
    descriptions = {
        'fecha': ('Clave temporal mensual unica; primer dia del mes', 'mes calendario'),
        'accidentes': ('Variable objetivo: total registrado', 'siniestros/mes'),
        **{key: ('Componente del total; descriptivo, excluido como predictor contemporaneo', 'siniestros/mes') for key in TYPES},
        'ano': ('Ano calendario', 'ano'), 'mes': ('Mes calendario', '1 a 12'),
        'dias_mes': ('Duracion del mes', 'dias'), 'tasa_diaria_mes': ('Total dividido por dias del mes; descriptiva', 'siniestros/dia'),
        'semestre': ('Semestre calendario', 'Primero/Segundo'),
        'particion': ('Papel temporal asignado antes de modelar', 'categoria'),
        'archivo': ('Documento original conservado', 'texto'), 'hoja': ('Hoja original', 'texto'),
        'celda_total': ('Celda del total mensual', 'referencia Excel'),
        'celdas_tipos': ('Rango de componentes', 'referencia Excel'), 'edicion': ('Anuario de publicacion', 'ano'),
    }
    dictionary = []
    for col in frame:
        description, unit = descriptions[col]
        values = frame[col]
        dictionary.append({'campo': col, 'descripcion': description, 'tipo': str(values.dtype),
                           'unidad': unit, 'faltantes_pct': 100*values.isna().mean(),
                           'cardinalidad': values.nunique(), 'minimo_o_ejemplo': str(values.min()),
                           'maximo_o_ejemplo': str(values.max())})
    save_table('diccionario_final.csv', pd.DataFrame(dictionary))
    raw_dict_rows=[
        dict(columna_original='B',significado='Ano o nombre del mes',tipo_original='texto',
             unidad='ano / mes calendario',rango='Enero a diciembre en cada bloque anual',
             faltantes_pct=0,guiones=0,observacion='Solo filas mensuales; no es identificador de accidente'),
        dict(columna_original='C',significado='Total de siniestros',tipo_original='entero',
             unidad='siniestros/mes',rango=f'{observed.accidentes.min()} a {observed.accidentes.max()}',
             faltantes_pct=0,guiones=0,observacion='Encabezados anuales separados como controles'),
    ]
    for i,name in enumerate(TYPES):
        originals=[v['valor_original'] for v in original if v['campo']==name]
        numbers=[v for v in originals if isinstance(v,(int,float))]
        raw_dict_rows.append(dict(columna_original=chr(68+i),significado=name,
                                 tipo_original='entero o guion',unidad='siniestros/mes',
                                 rango=f'{min(numbers)} a {max(numbers)}; guion si corresponde',
                                 faltantes_pct=0,guiones=originals.count('-'),
                                 observacion='Guion interpretado como cero solo tras validar sumas'))
    raw_dictionary=pd.DataFrame(raw_dict_rows)
    save_table('diccionario_original.csv', raw_dictionary)
    audit = dict(archivos=len(inventory), filas_extraidas=len(observed), filas_finales=len(frame),
                 duplicados_coincidentes=len(observed)-len(frame), guiones_convertidos=conversions,
                 meses_faltantes=0, objetivos_imputados=0, sumas_inconsistentes=0,
                 totales_anuales_validados=len(annual), distribucion_particiones=frame['particion'].value_counts().to_dict(),
                 fuente='INE; Patrulla Caminera, zona Central e Interior',
                 limite='Coherencia verificada; estabilidad institucional de cobertura no demostrada')
    save_json('auditoria_datos.json', audit)
    print('Fase 1 - preparacion:', audit)
    return frame


if __name__ == '__main__':
    main()
