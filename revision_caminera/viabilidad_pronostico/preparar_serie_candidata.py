"""Lee cuadros INE sin modificarlos y prepara evidencia de viabilidad (JSON).

No ajusta modelos ni convierte frecuencias en personas. El formato XLSX de la
fuente se conserva y cada valor mensual incluye su hoja y celda originales.
"""
import collections
import json
import re
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent
MONTHS = dict(enero=1, febrero=2, marzo=3, abril=4, mayo=5, junio=6,
              julio=7, agosto=8, setiembre=9, septiembre=9, octubre=10,
              noviembre=11, diciembre=12)
TYPES = ['choque', 'roce', 'vuelco', 'arrollamiento', 'incendio', 'casual', 'atropellamiento']


def number(value):
    if value == '-':
        return 0
    if isinstance(value, (int, float)) and value >= 0 and int(value) == value:
        return int(value)
    raise ValueError(f'Conteo no valido: {value!r}')


def main():
    candidates = collections.defaultdict(list)
    annual = []
    checks = []
    for file in sorted((ROOT / 'fuentes').glob('*tipologia*.xlsx')):
        workbook = load_workbook(file, read_only=True, data_only=True)
        sheet = workbook.worksheets[0]
        year = None
        for row_index, row in enumerate(sheet.iter_rows(values_only=True), 1):
            label = str(row[1] or '').strip().lower() if len(row) > 1 else ''
            match = re.fullmatch(r'año\s+(20\d{2})', label)
            if match:
                year = int(match.group(1))
                annual.append(dict(year=year, total=number(row[2]), source=file.name,
                                   sheet=sheet.title, cell=f'C{row_index}'))
            elif year and label in MONTHS:
                total = number(row[2])
                record = dict(fecha=f'{year}-{MONTHS[label]:02d}', accidentes=total,
                              **{name: number(row[3 + i]) for i, name in enumerate(TYPES)})
                record.update(source=file.name, sheet=sheet.title, cell=f'C{row_index}',
                              cell_types=f'D{row_index}:J{row_index}')
                candidates[record['fecha']].append(record)
                if sum(record[name] for name in TYPES) != total:
                    checks.append(dict(tipo='tipos_no_suman_total', record=record))
        workbook.close()

    selected = []
    for date, records in sorted(candidates.items()):
        signatures = {tuple(record[key] for key in ['accidentes'] + TYPES) for record in records}
        if len(signatures) > 1:
            checks.append(dict(tipo='discrepancia_entre_ediciones', fecha=date, registros=records))
        # Las ediciones mas recientes se ordenan al final del nombre de archivo.
        selected.append({**records[-1], 'otras_fuentes': records[:-1]})
    for a in annual:
        computed = sum(r['accidentes'] for r in selected if int(r['fecha'][:4]) == a['year'])
        if computed != a['total']:
            checks.append(dict(tipo='suma_meses_diferente_total_publicado', publicado=a, suma_meses=computed))

    per_year = collections.defaultdict(int)
    for record in selected:
        per_year[record['fecha'][:4]] += record['accidentes']
    result = dict(unidad='Siniestros registrados por la Patrulla Caminera en zona Central e Interior, por mes',
                  estado='Serie candidata para diseno del proyecto; comparabilidad institucional pendiente de documentar',
                  inicio=selected[0]['fecha'], fin=selected[-1]['fecha'], meses=len(selected),
                  totales_anuales=dict(per_year), incidencias=checks, registros=selected)
    (ROOT / 'serie_mensual_candidata_2016_2024.json').write_text(
        json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps({key: value for key, value in result.items() if key != 'registros'}, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
